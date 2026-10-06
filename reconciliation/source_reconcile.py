"""Independent original-PDF -> workbook -> monthly reconciliation.

Run from repository root: python reconciliation/source_reconcile.py
PDF coordinates are specific to the supplied, hashed original reports. They
were checked against rendered pages; no workbook values drive PDF extraction.
"""
from pathlib import Path
from collections import defaultdict
from datetime import date
import csv
import hashlib
import json
import sys

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
SOURCE = ROOT / 'reconciliation' / 'sources'
if not SOURCE.exists():
    SOURCE = ROOT.parent
OUT = ROOT / 'reconciliation'


def grid_rows(page, bounds, top, bottom):
    groups = defaultdict(list)
    for char in page.chars:
        if top <= char['top'] <= bottom and char['text'].strip():
            groups[round(char['top'], 1)].append(char)
    rows = []
    for y, chars in sorted(groups.items()):
        cells = [''.join(c['text'] for c in sorted(chars, key=lambda c: c['x0'])
                         if left <= (c['x0'] + c['x1']) / 2 < right)
                 for left, right in zip(bounds, bounds[1:])]
        if cells[0]:
            rows.append((y, cells))
    return rows


def number(value):
    if value is None or str(value).strip() in ('', '-'):
        return None
    return int(str(value).replace(',', '').replace(' ', ''))


def write_csv(name, rows):
    with (OUT / name).open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    files = {
        'workbook': SOURCE / 'Antipolo_Disease_Surveillance_2016-2025.xlsx',
        'Dengue': next(SOURCE.glob('Distribution of Dengue*.pdf')),
        'Measles-Rubella': next(SOURCE.glob('Distribution of Measles-Rubella*.pdf')),
        'Leptospirosis': SOURCE / 'LEPTO 2016-2025 ANTIPOLO.pdf',
    }
    report = {'sources': {key: {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
                          for key, path in files.items()}}
    pdf_values, totals = {}, {}
    for disease in ('Leptospirosis', 'Measles-Rubella', 'Dengue'):
        with pdfplumber.open(files[disease]) as pdf:
            text = '\n\n'.join(f'PAGE {i + 1}\n{p.extract_text(x_tolerance=5, y_tolerance=2)}'
                               for i, p in enumerate(pdf.pages))
            (OUT / f'source-{disease.lower()}-text.txt').write_text(text, encoding='utf-8')
            if disease == 'Leptospirosis':
                rows = [(1, None, r) for t in pdf.pages[0].extract_tables() for r in t
                        if r and str(r[0]).isdigit() and 1 <= int(r[0]) <= 53]
                years = list(range(2016, 2026))
            elif disease == 'Measles-Rubella':
                bounds = [49, 105] + [105 + i * 42.6 for i in range(1, 11)]
                rows = [(1, y, r) for y, r in grid_rows(pdf.pages[0], bounds, 190, 725)]
                years = list(range(2025, 2015, -1))
            else:
                rows = [(1, None, r) for t in pdf.pages[0].extract_tables() for r in t]
                bounds = [106, 179] + [179 + i * 31.9 for i in range(1, 11)]
                rows += [(2, y, r) for y, r in grid_rows(pdf.pages[1], bounds, 167, 501)]
                years = list(range(2016, 2026))
            for page, y, row in rows:
                if row[0].replace(' ', '') in ('TOTAL', 'GRANDTOTAL'):
                    for year, value in zip(years, row[1:]):
                        totals[disease, year] = number(value)
                    continue
                week = int(row[0])
                assert len(row) == 11
                for year, value in zip(years, row[1:]):
                    key = disease, year, week
                    assert key not in pdf_values
                    pdf_values[key] = {'value': number(value), 'pdf_cell': value, 'page': page, 'pdf_top': y}
        assert sum(k[0] == disease for k in pdf_values) == 530

    workbook = openpyxl.load_workbook(files['workbook'], data_only=False)
    cached = openpyxl.load_workbook(files['workbook'], data_only=True)
    cells, annual, monthly = [], [], defaultdict(int)
    for disease in ('Leptospirosis', 'Measles-Rubella', 'Dengue'):
        sheet = workbook[disease]
        for col, year in enumerate(range(2016, 2026), 2):
            values = []
            for week in range(1, 54):
                cell = sheet.cell(week + 3, col)
                val = number(cell.value)
                pdf = pdf_values[disease, year, week]
                values.append(val)
                cells.append(dict(disease=disease, year=year, week=week, workbook_cell=cell.coordinate,
                                  workbook_value=val, pdf_value=pdf['value'], pdf_cell=pdf['pdf_cell'],
                                  pdf_page=pdf['page'], pdf_top=pdf['pdf_top'], matches=val == pdf['value']))
                if val is not None:
                    try:
                        thursday = date.fromisocalendar(year, week, 4)
                        calendar_year, month = thursday.year, thursday.month
                        fallback = False
                    except ValueError:
                        assert week == 53
                        calendar_year, month, fallback = year, 12, True
                    monthly[disease, calendar_year, month] += val
                    cells[-1].update(mapped_year=calendar_year, mapped_month=month, week53_fallback=fallback)
                else:
                    cells[-1].update(mapped_year=None, mapped_month=None, week53_fallback=None)
            total = sum(v for v in values if v is not None)
            printed = totals.get((disease, year))
            annual.append(dict(disease=disease, year=year, pdf_printed_total=printed,
                               pdf_weekly_sum=sum(pdf_values[disease, year, w]['value'] or 0 for w in range(1,54)),
                               workbook_weekly_sum=total,
                               workbook_total_formula=sheet.cell(57, col).value,
                               workbook_total_cached=cached[disease].cell(57,col).value,
                               weekly_minus_printed=None if printed is None else total-printed,
                               blank_cells=sum(v is None for v in values),
                               explicit_zero_cells=values.count(0), week53=values[-1],
                               status='No printed total; weekly cells match' if printed is None else
                                      'Printed/weekly discrepancy; CHO correction required' if printed != total else
                                      'Printed total and weekly cells reconcile',
                               follow_up='CHO epidemiological completeness, age, classification and week calendar verification required'))
    assert all(row['matches'] for row in cells), [r for r in cells if not r['matches']]
    from dashboard.data.xlsx_parser import parse_surveillance_xlsx
    parsed, notes = parse_surveillance_xlsx(files['workbook'].read_bytes())
    # The legacy parser renames the combined source category; keep this auditable.
    aliases = {'Measles': 'Measles-Rubella'}
    parsed_map = {(aliases.get(r.disease, r.disease), int(r.year), int(r.month)): int(r.cases)
                  for r in parsed.itertuples()}
    assert parsed_map == dict(monthly)
    report.update(weekly_cells_compared=len(cells), matching_weekly_cells=sum(r['matches'] for r in cells),
                  numeric_cells=sum(r['workbook_value'] is not None for r in cells),
                  blank_cells=sum(r['workbook_value'] is None for r in cells),
                  explicit_zero_cells=sum(r['workbook_value'] == 0 for r in cells),
                  monthly_rows=len(monthly), monthly_rows_match_parser=True, parser_notes=notes,
                  parser_labels=sorted(parsed.disease.unique().tolist()),
                  annual=annual, workbook_notes=[r[0] for r in workbook['Notes'].values if r[0]])
    write_csv('source-weekly-cell-comparison.csv', cells)
    write_csv('source-annual-reconciliation.csv', annual)
    write_csv('source-monthly-independent.csv', [dict(disease=d, year=y, month=m, cases=v)
                                                for (d,y,m),v in sorted(monthly.items())])
    (OUT / 'source-reconciliation.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('sources','annual','workbook_notes')}, indent=2))


if __name__ == '__main__':
    main()
